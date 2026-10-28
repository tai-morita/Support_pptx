"""PowerPointへ通常画像、注釈、Zoomを描画する。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

from matching import MatchedImages
from models import CmPoint, CmSize, DataSetConfig, FolderConfig, ZoomConfig


ANNOTATION_FONT_SIZE_PT = 12.0
"""!< 画像注釈のフォントサイズ。"""
ANNOTATION_GAP_CM = 0.15
"""!< 通常画像と注釈の間隔。"""
ANNOTATION_HEIGHT_CM = 0.8
"""!< 注釈テキストボックスの高さ。"""
ANNOTATION_BOTTOM_MARGIN_CM = 0.4
"""!< 注釈とスライド下端の余白。"""
TITLE_IMAGE_GAP_CM = 0.2
"""!< タイトルと通常画像領域の間隔。"""
MIN_HORIZONTAL_GAP_CM = 0.3
"""!< 通常画像間および左右端の最小余白。"""


@dataclass(frozen=True)
class ImagePlacement:
    """通常画像の配置位置と大きさを保持する。"""

    location: CmPoint
    """!< 画像を貼り付ける左上座標。"""
    size: CmSize
    """!< 画像を貼り付ける大きさ。"""


def add_blank_slide(presentation: Any) -> Any:
    """! プレゼンテーション末尾へ空白スライドを追加する。

    @param presentation 追加先のPresentation。
    @return 追加したSlide。
    """
    layout_index = 6 if len(presentation.slide_layouts) > 6 else -1
    return presentation.slides.add_slide(
        presentation.slide_layouts[layout_index]
    )


def add_picture(
    slide: Any,
    image_path: Path,
    location: CmPoint,
    size: CmSize,
) -> Any:
    """! 指定位置と大きさで画像を追加する。"""
    return slide.shapes.add_picture(
        str(image_path),
        Cm(location.left),
        Cm(location.top),
        width=Cm(size.width),
        height=Cm(size.height),
    )


def add_centered_text(
    slide: Any,
    text: str,
    location: CmPoint,
    size: CmSize,
    font_size_pt: float,
) -> Any:
    """! 枠線と塗りつぶしのない中央揃えテキストを追加する。"""
    textbox = slide.shapes.add_textbox(
        Cm(location.left),
        Cm(location.top),
        Cm(size.width),
        Cm(size.height),
    )
    textbox.fill.background()
    textbox.line.fill.background()
    text_frame = textbox.text_frame
    text_frame.clear()
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.text = text.replace("\r\n", "\n").replace("\r", "\n")
    paragraph.font.size = Pt(font_size_pt)
    for run in paragraph.runs:
        run.font.size = Pt(font_size_pt)
    return textbox


def add_title(
    slide: Any,
    title: str,
    location: CmPoint,
    size: CmSize,
    font_size_pt: float,
) -> Any:
    """! スライド上部へ一致条件タイトルを追加する。"""
    return add_centered_text(
        slide,
        title,
        location,
        size,
        font_size_pt,
    )


def add_annotation(
    slide: Any,
    folder: FolderConfig,
    placement: ImagePlacement,
    top: float,
) -> Any:
    """! 通常画像と同じ横位置でスライド下部へ注釈を追加する。"""
    location = CmPoint(
        left=placement.location.left,
        top=top,
    )
    return add_centered_text(
        slide,
        folder.annotation,
        location,
        CmSize(width=placement.size.width, height=ANNOTATION_HEIGHT_CM),
        ANNOTATION_FONT_SIZE_PT,
    )


def crop_ratios(
    source_size: CmSize,
    source_location: CmPoint,
    zoom: ZoomConfig,
) -> tuple[float, ...]:
    """! スライド絶対cmの切り取り矩形をpython-pptxの比率へ変換する。"""
    cut_left = zoom.cut_location.left - source_location.left
    cut_top = zoom.cut_location.top - source_location.top
    right_edge = cut_left + zoom.cut_size.width
    bottom_edge = cut_top + zoom.cut_size.height
    if (
        cut_left < 0
        or cut_top < 0
        or right_edge > source_size.width
        or bottom_edge > source_size.height
    ):
        exceeded_edges = []
        if cut_left < 0:
            exceeded_edges.append(
                f"左端が通常画像の左端より {abs(cut_left):.3f} cm 左"
            )
        if cut_top < 0:
            exceeded_edges.append(
                f"上端が通常画像の上端より {abs(cut_top):.3f} cm 上"
            )
        if right_edge > source_size.width:
            exceeded_edges.append(
                f"右端={right_edge:.3f} cm > 幅={source_size.width:.3f} cm"
            )
        if bottom_edge > source_size.height:
            exceeded_edges.append(
                f"下端={bottom_edge:.3f} cm > 高さ={source_size.height:.3f} cm"
            )
        raise ValueError(
            f"Zoom '{zoom.name}' の切り取り範囲が通常画像からはみ出します: "
            f"cut_location(スライド座標)=[横={zoom.cut_location.left:.3f}, "
            f"縦={zoom.cut_location.top:.3f}] cm, "
            f"通常画像位置=[横={source_location.left:.3f}, "
            f"縦={source_location.top:.3f}] cm, "
            f"cut_size=[高さ={zoom.cut_size.height:.3f}, "
            f"幅={zoom.cut_size.width:.3f}] cm; "
            f"切り取り位置(通常画像内)=[横={cut_left:.3f}, "
            f"縦={cut_top:.3f}] cm; "
            f"{', '.join(exceeded_edges)} "
            f"(通常画像サイズ: 幅={source_size.width:.3f}, "
            f"高さ={source_size.height:.3f} cm)"
        )
    return (
        cut_left / source_size.width,
        cut_top / source_size.height,
        1.0 - (right_edge / source_size.width),
        1.0 - (bottom_edge / source_size.height),
    )


def add_zoom(
    slide: Any,
    image_path: Path,
    source_size: CmSize,
    source_location: CmPoint,
    zoom: ZoomConfig,
) -> Any:
    """! 元画像の指定範囲を切り取り、拡大画像として追加する。"""
    picture = add_picture(slide, image_path, zoom.location, zoom.size)
    crop_left, crop_top, crop_right, crop_bottom = crop_ratios(
        source_size,
        source_location,
        zoom,
    )
    picture.crop_left = crop_left
    picture.crop_top = crop_top
    picture.crop_right = crop_right
    picture.crop_bottom = crop_bottom
    return picture


def _image_aspect_ratio(image_path: Path) -> float:
    """! 画像の先頭フレームから幅÷高さを取得する。"""
    try:
        with Image.open(image_path) as image:
            width, height = image.size
    except (OSError, ValueError, SyntaxError) as exc:
        raise ValueError(f"画像サイズを取得できません: {image_path}") from exc
    if width <= 0 or height <= 0:
        raise ValueError(f"画像サイズが不正です: {image_path}")
    return width / height


def annotation_top(presentation: Any) -> float:
    """! スライド下部に固定する注釈領域の上端を返す。"""
    slide_height_cm = presentation.slide_height / Cm(1)
    return (
        slide_height_cm
        - ANNOTATION_BOTTOM_MARGIN_CM
        - ANNOTATION_HEIGHT_CM
    )


def calculate_image_placements(
    presentation: Any,
    dataset: DataSetConfig,
    image_paths: tuple[Path, ...],
) -> tuple[ImagePlacement, ...]:
    """! 元画像比率を維持して通常画像の位置と大きさを計算する。

    @param presentation スライド寸法を取得するPresentation。
    @param dataset タイトル領域を含むDataSet設定。
    @param image_paths Folder設定と同じ順序の画像パス。
    @return 入力順に並べた画像配置。
    """
    image_count = len(image_paths)
    if image_count != len(dataset.folders) or image_count < 2:
        raise ValueError("画像数はFolder数と同じ2件以上にしてください")

    slide_width_cm = presentation.slide_width / Cm(1)
    image_top = (
        dataset.slide_title_location.top
        + dataset.slide_title_size.height
        + TITLE_IMAGE_GAP_CM
    )
    image_bottom = annotation_top(presentation) - ANNOTATION_GAP_CM
    available_height = image_bottom - image_top
    if available_height <= 0:
        raise ValueError("タイトルと注釈を除いた画像領域の高さがありません")

    aspect_ratios = tuple(_image_aspect_ratio(path) for path in image_paths)
    horizontal_gap_count = image_count + 1
    available_width = (
        slide_width_cm - MIN_HORIZONTAL_GAP_CM * horizontal_gap_count
    )
    if available_width <= 0:
        raise ValueError("画像を配置できるスライド横幅がありません")

    height_by_width = available_width / sum(aspect_ratios)
    image_height = min(available_height, height_by_width)
    if image_height <= 0:
        raise ValueError("通常画像の自動計算サイズが不正です")

    widths = tuple(ratio * image_height for ratio in aspect_ratios)
    horizontal_gap = (
        slide_width_cm - sum(widths)
    ) / horizontal_gap_count
    image_y = image_top + (available_height - image_height) / 2.0
    placements: list[ImagePlacement] = []
    image_x = horizontal_gap
    for width in widths:
        placements.append(
            ImagePlacement(
                location=CmPoint(left=image_x, top=image_y),
                size=CmSize(width=width, height=image_height),
            )
        )
        image_x += width + horizontal_gap
    return tuple(placements)


def render_comparison_slide(
    presentation: Any,
    dataset: DataSetConfig,
    matched: MatchedImages,
    zoom_name: str | None = None,
) -> Any:
    """! 1条件分の通常画像と任意の同名Zoomを1枚へ描画する。

    @param presentation 描画先のPresentation。
    @param dataset Folder配置を含むDataSet設定。
    @param matched Folder順の一致画像群。
    @param zoom_name 重ねるZoom名。通常スライドではNone。
    @return 描画したSlide。
    """
    slide = add_blank_slide(presentation)
    placements = calculate_image_placements(
        presentation,
        dataset,
        matched.paths,
    )
    annotation_y = annotation_top(presentation)
    add_title(
        slide,
        matched.title,
        dataset.slide_title_location,
        dataset.slide_title_size,
        dataset.slide_title_font_size_pt,
    )
    for folder, image_path, placement in zip(
        dataset.folders,
        matched.paths,
        placements,
    ):
        add_picture(slide, image_path, placement.location, placement.size)
        add_annotation(slide, folder, placement, annotation_y)

    if zoom_name is not None:
        for folder, image_path, placement in zip(
            dataset.folders,
            matched.paths,
            placements,
        ):
            zoom = next(
                item for item in folder.zooms if item.name == zoom_name
            )
            try:
                add_zoom(
                    slide,
                    image_path,
                    placement.size,
                    placement.location,
                    zoom,
                )
            except ValueError as exc:
                raise ValueError(
                    f"Folder '{folder.annotation}': {exc}"
                ) from exc
    return slide


def load_presentation(input_pptx: Path) -> Any:
    """! 入力PowerPointを読み込む。"""
    return Presentation(str(input_pptx))
